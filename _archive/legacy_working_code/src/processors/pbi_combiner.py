"""
Power BI raw file combiner -- PBI Combine Tool V1.

Reads the raw .xlsx files produced by PBI Download Tool V1
(src/collectors/pbi_collector.py -- NOT imported or modified here; the
two tools are independent) and combines them into one Excel file per
report per requested date range.

Combine rule (reverse-engineered from hand-made reference files and
verified to match them exactly, for all 3 reports):
  - Every raw export ends with exactly 3 rows that are NOT data:
    a repeated "Total" row, a blank row, and a filter-description note
    row. These are dropped structurally (by position from the end),
    not by matching their (garbled-encoding-prone) text.
  - The real header lives in row 1 of the raw file. Some reports have
    one or more additional header rows immediately below it (e.g. "WD
    WL by Group" has a second row of column-group labels); those extra
    header rows are dropped entirely, not merged into anything.
  - In the combined output, row 1's column A is blanked out (it held
    "Brand"/"Withdrawal Group" in the raw file, which no longer means
    anything once dates are mixed together).
  - Every remaining data row's column A (which the raw file left mostly
    blank, or set to "Vantage" for the top summary row) is replaced
    with the file's source_date, taken only from the filename -- never
    from file metadata or download time.
  - Rows are appended in ascending date order, in whatever row order
    they already had in each day's raw file. Header appears once.
  - Column alignment is by column NAME, not raw position: some days'
    export genuinely omits a column (observed for "WD WL by Group" --
    Power BI appears to drop a column entirely, rather than keep it
    with blank values, when that column has zero data for the day).
    The canonical column set for a report+range is the longest header
    seen among the files being combined; any file whose header is a
    same-order subset of that canonical header gets its row values
    aligned to the canonical positions, with missing columns filled
    blank. A file whose header is NOT a same-order subset (an actual
    extra/reordered column, not just a missing one) is a genuine
    schema error and blocks output, same as before.
"""

import re
from datetime import date as _date, timedelta
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[2]
RAW_ROOT = PROJECT_ROOT / "data" / "raw"
COMBINE_ROOT = PROJECT_ROOT / "data" / "COMBINE"

# Rows that are always non-data at the tail of a raw export, regardless
# of report: [repeated Total row, blank row, filter-description note].
TRAILING_NON_DATA_ROWS = 3

REPORT_SPECS = {
    "op_workload_p1": {
        "display_name": "OP Workload P.1",
        "raw_subfolder": "OP Workload P.1",
        "raw_prefix": "OP_Workload_P1",
        # NOTE: double underscore after P1 is intentional -- matches
        # the user-provided reference file naming exactly.
        "output_name_template": "COMBINE-OP_Workload_P1__{start} to {end}.xlsx",
        "header_rows_in_raw": 1,
    },
    "op_workload_p2": {
        "display_name": "OP Workload P.2",
        "raw_subfolder": "OP Workload P.2",
        "raw_prefix": "OP_Workload_P2",
        "output_name_template": "COMBINE-OP_Workload_P2_{start} to {end}.xlsx",
        "header_rows_in_raw": 1,
    },
    "wd_wl_by_group": {
        "display_name": "WD WL by Group",
        "raw_subfolder": "WD WL by Group",
        "raw_prefix": "WD_WL_by_Group",
        "output_name_template": "COMBINE-WD_WL_by_Group_{start} to {end}.xlsx",
        "header_rows_in_raw": 2,
    },
}


def _daterange(start: _date, end: _date):
    d = start
    while d <= end:
        yield d
        d += timedelta(days=1)


def _scan_raw_files(spec: dict):
    """Scan the report's raw folder; parse source_date from filename only.

    Returns (by_date: {date: Path}, filename_errors: [filename]).
    """
    folder = RAW_ROOT / spec["raw_subfolder"]
    by_date = {}
    filename_errors = []
    if not folder.exists():
        return by_date, filename_errors

    pattern = re.compile(rf"^{re.escape(spec['raw_prefix'])}_(\d{{4}}-\d{{2}}-\d{{2}})\.xlsx$")
    for f in folder.glob(f"{spec['raw_prefix']}_*.xlsx"):
        if f.name.startswith("~$"):
            continue
        m = pattern.match(f.name)
        if not m:
            filename_errors.append(f.name)
            continue
        try:
            d = _date.fromisoformat(m.group(1))
        except ValueError:
            filename_errors.append(f.name)
            continue
        by_date[d] = f
    return by_date, filename_errors


def _read_raw(path: Path, header_rows_in_raw: int):
    """Read one raw export. Returns (header_row_values, data_rows)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]

    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

    data_start = header_rows_in_raw + 1
    data_end = ws.max_row - TRAILING_NON_DATA_ROWS
    rows = []
    for r in range(data_start, data_end + 1):
        rows.append([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])

    wb.close()
    return header, rows


def _build_column_mapping(file_header: list, canonical_header: list):
    """
    Map each canonical column position to the matching position in
    file_header (or None if file_header doesn't have that column),
    requiring file_header's columns to appear in canonical_header in
    the same relative order (a "subsequence" match).

    Raises ValueError if file_header has a column not present in
    canonical_header, or columns out of order -- a genuine schema
    conflict, not just a missing optional column.
    """
    mapping = []
    file_idx = 0
    for canon_col in canonical_header:
        if file_idx < len(file_header) and file_header[file_idx] == canon_col:
            mapping.append(file_idx)
            file_idx += 1
        else:
            mapping.append(None)
    if file_idx != len(file_header):
        raise ValueError("file_header is not a same-order subset of canonical_header")
    return mapping


def _write_combined(output_path: Path, header: list, rows: list):
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Export"

    ws.append(header)
    for row in rows:
        ws.append(row)

    for r in range(2, len(rows) + 2):
        ws.cell(row=r, column=1).number_format = "mm-dd-yy"

    wb.save(output_path)
    wb.close()


def combine_report(report: str, start_date: str, end_date: str) -> dict:
    """
    PBI Combine Tool V1 -- per-report entry point.

    Args:
        report: one of "op_workload_p1", "op_workload_p2", "wd_wl_by_group"
        start_date: "YYYY-MM-DD", inclusive
        end_date: "YYYY-MM-DD", inclusive

    Returns a result dict (also see combine_date_range). No output file
    is written unless status is COMPLETE.
    """
    if report not in REPORT_SPECS:
        raise ValueError(f"Unknown report '{report}'. Valid: {list(REPORT_SPECS)}")
    spec = REPORT_SPECS[report]

    start = _date.fromisoformat(start_date)
    end = _date.fromisoformat(end_date)
    expected_dates = list(_daterange(start, end))

    by_date, filename_errors = _scan_raw_files(spec)

    missing = [d.isoformat() for d in expected_dates if d not in by_date]
    found_dates = [d for d in expected_dates if d in by_date]

    result = {
        "report": spec["display_name"],
        "report_key": report,
        "start_date": start_date,
        "end_date": end_date,
        "expected": len(expected_dates),
        "found": len(found_dates),
        "missing": missing,
        "filename_errors": filename_errors,
        "schema_errors": [],
        "rows": 0,
        "output_path": None,
        "status": "INCOMPLETE",
    }

    if missing:
        result["status"] = "INCOMPLETE"
        return result
    if filename_errors:
        result["status"] = "FAILED"
        return result

    loaded = {}
    headers_by_date = {}
    for d in found_dates:
        header, rows = _read_raw(by_date[d], spec["header_rows_in_raw"])
        headers_by_date[d] = header[1:]
        loaded[d] = rows

    # Canonical column set = the longest header seen. If more than one
    # distinct header shares that max length, that's ambiguous -- treat
    # every date not matching the first-seen longest header as an error.
    max_len = max(len(h) for h in headers_by_date.values())
    canonical = None
    for d in found_dates:
        if len(headers_by_date[d]) == max_len:
            canonical = headers_by_date[d]
            break

    combined_rows = []
    for d in found_dates:
        header = headers_by_date[d]
        if header == canonical:
            mapping = list(range(len(canonical)))
        else:
            try:
                mapping = _build_column_mapping(header, canonical)
            except ValueError:
                result["schema_errors"].append({
                    "date": d.isoformat(),
                    "file": by_date[d].name,
                    "expected_columns": canonical,
                    "found_columns": header,
                })
                continue

        for row in loaded[d]:
            row_wo_a = row[1:]
            aligned = [row_wo_a[idx] if idx is not None else None for idx in mapping]
            combined_rows.append([d] + aligned)

    if result["schema_errors"]:
        result["status"] = "FAILED"
        return result

    out_header = [None] + list(canonical)

    output_name = spec["output_name_template"].format(start=start_date, end=end_date)
    output_path = COMBINE_ROOT / output_name
    _write_combined(output_path, out_header, combined_rows)

    result["rows"] = len(combined_rows)
    result["output_path"] = str(output_path)
    result["status"] = "COMPLETE"
    return result


def combine_date_range(start_date: str, end_date: str, reports=None) -> dict:
    """
    PBI Combine Tool V1 -- official high-level entry point.

    Args:
        start_date, end_date: "YYYY-MM-DD", inclusive
        reports: list of report keys to combine; None = all 3

    Returns {report_key: result_dict}. Prints a human-readable summary.
    """
    keys = reports if reports else list(REPORT_SPECS.keys())
    for key in keys:
        if key not in REPORT_SPECS:
            raise ValueError(f"Unknown report '{key}'. Valid: {list(REPORT_SPECS)}")

    results = {key: combine_report(key, start_date, end_date) for key in keys}
    _print_summary(start_date, end_date, results)
    return results


def _print_summary(start_date: str, end_date: str, results: dict) -> None:
    print("PBI Combine Completed")
    print()
    print("Date Range:")
    print(f"{start_date} -> {end_date}")
    print()
    for result in results.values():
        print(result["report"])
        print(f"Expected: {result['expected']}")
        print(f"Found: {result['found']}")
        if result["missing"]:
            print("Missing:")
            for d in result["missing"]:
                print(f"  {d}")
        if result["filename_errors"]:
            print("Filename Errors:")
            for f in result["filename_errors"]:
                print(f"  {f}")
        if result["schema_errors"]:
            print("Schema Errors:")
            for e in result["schema_errors"]:
                print(f"  {e['date']} ({e['file']})")
                print(f"    expected columns: {e['expected_columns']}")
                print(f"    found columns:    {e['found_columns']}")
        print(f"Rows: {result['rows']}")
        print(f"Status: {result['status']}")
        print()

    outputs = [r["output_path"] for r in results.values() if r["output_path"]]
    if outputs:
        print("Outputs:")
        for o in outputs:
            print(o)
