"""
Lark Chatgroup Volumn combiner.

Reads a manually-downloaded Lark chatgroup export file from
data/raw/lark_chatgroup/Group Management.xlsx (the user re-downloads
and overwrites this file by hand -- there is no automated Lark
download yet) and produces two consolidated Excel files per requested
date range:

  - "AC Chatgroup Volumn <start> to <end>.xlsx"
    from every sheet in Group Management.xlsx EXCEPT the 3 DW sheets
    listed below (see TARGETS["ac"]["exclude_sheet_names"]), combined
    into one output sheet named "AC Chatgroup Volumn". OP column
    matched by header "OP Name (Actual)" (case-insensitive; also
    matches sheets that spell it "OP name (Actual)").

  - "DW Chatgroup Volumn <start> to <end>.xlsx"
    from 3 specific sheets in the SAME Group Management.xlsx file:
    "Deposit Support", "Withdrawal Support", "Deposit & Withdrawal
    Support". OP column matched by header "OP name".

  Confirmed with the user 2026-09-21 that this single-merged-file
  layout is the standing format going forward -- an earlier version of
  this tool read AC and DW from 4 separate files (AC from its own
  Group Management.xlsx, DW from 3 separate "... Support (1).xlsx"
  files each with a "Group chat requests" sheet); those 3 DW files
  have since been deleted from data/raw/lark_chatgroup/.

Both outputs follow the same rule:
  - keep a row only if its Date falls within [start_date, end_date]
    inclusive
  - keep a row only if the OP-name column is not blank
  - header row once, then all kept rows in the order encountered
    (source file/sheet by source file/sheet, top to bottom)

Lark Name -> OP Name transfer (applied to every kept row before the
file is written -- there is no separate "raw" output anymore, this IS
the final Combine output):
  - data/raw/roster/Roster.xlsx (also manually re-downloaded by the
    user for now) maps each person's "Lark Name" (the display name
    Lark chat data uses, which is what ends up in the OP-name column
    collected above) to their canonical "OP Name". The two often
    differ in spelling/word order for the same person.
  - if the row's OP-name value matches a Roster "Lark Name" entry, the
    OP Name column in the output is REPLACED with Roster's "OP Name"
    for that person, and column E ("Roster Status") is left blank.
  - if there is no match (most commonly because the person has left
    and been removed from the Roster), the OP Name column is left
    UNCHANGED (not converted), and column E is set to
    "Inactive from Roster" so these rows are easy to spot/filter.
  - matching is done on the Lark Name string with leading/trailing
    whitespace stripped, case-sensitive otherwise.

Output columns: A = Date (date-only, YYYY/MM/DD), B = OP Name (after
Roster transfer), C = Group Name, D = Source (the Group Management.xlsx
sheet name the row came from, for both "ac" and "dw"), E = Roster
Status ("Inactive from Roster" or blank).

Column matching (Date/Group Name/OP-name/Lark Name/OP Name headers) is
by HEADER NAME (case-insensitive), not by column position -- the
source sheets do not all use the same column order, and some omit
columns others have. This was verified against the actual source
files, not assumed from a description.

Known source-file quirk worked around here: some of these files have
wrong/stale worksheet dimension metadata (openpyxl under-reports
max_row/max_column for them), so row/column ranges are always read
with an explicit generous bound rather than trusting that metadata.
"""

import re
from datetime import date as _date, datetime
from pathlib import Path

import openpyxl

PROJECT_ROOT = Path(__file__).resolve().parents[2]
SOURCE_DIR = PROJECT_ROOT / "data" / "raw" / "lark_chatgroup"
ROSTER_PATH = PROJECT_ROOT / "data" / "raw" / "roster" / "Roster.xlsx"
OUTPUT_DIR = PROJECT_ROOT / "data" / "COMBINE"

INACTIVE_FLAG = "Inactive from Roster"

# Generous fixed bounds used for every sheet read, because some source
# files report incorrect (too-small) worksheet dimensions.
MAX_COL_BOUND = 30
MAX_ROW_BOUND = 200000

#
# As of 2026-09-21, the source is a single merged "Group Management.xlsx"
# with all AC and DW sheets inside one workbook -- confirmed with the
# user this is now the standing format going forward (the old layout,
# AC in its own "Group Management.xlsx" + DW split across 3 separate
# files each with a "Group chat requests" sheet, is retired; those 3
# files have been deleted from data/raw/lark_chatgroup/). Both targets
# read from this one file; the split between them is by sheet name:
# DW = the 3 sheets explicitly listed for "dw" below; AC = every other
# sheet in the file (via "exclude_sheet_names").
TARGETS = {
    "ac": {
        "display_name": "AC Chatgroup Volumn",
        "source_files": ["Group Management.xlsx"],
        "sheet_names": None,  # None = every sheet in the file except exclude_sheet_names
        "exclude_sheet_names": ["Deposit Support", "Withdrawal Support", "Deposit & Withdrawal Support"],
        "op_col_candidates": ["OP Name (Actual)"],
        "output_name_template": "AC Chatgroup Volumn {start} to {end}.xlsx",
        "source_label_mode": "sheet_name",
    },
    "dw": {
        "display_name": "DW Chatgroup Volumn",
        "source_files": ["Group Management.xlsx"],
        "sheet_names": ["Deposit Support", "Withdrawal Support", "Deposit & Withdrawal Support"],
        "exclude_sheet_names": None,
        "op_col_candidates": ["OP name"],
        "output_name_template": "DW Chatgroup Volumn {start} to {end}.xlsx",
        "source_label_mode": "sheet_name",
    },
}


def _clean_file_label(filename: str) -> str:
    """'Deposit & Withdrawal Support (1).xlsx' -> 'Deposit & Withdrawal Support'"""
    stem = Path(filename).stem
    return re.sub(r"\s*\(\d+\)\s*$", "", stem).strip()


def _find_col(header_row, *candidates):
    normalized = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        normalized.setdefault(str(h).strip().lower(), i)
    for cand in candidates:
        idx = normalized.get(cand.strip().lower())
        if idx is not None:
            return idx
    return None


def _collect_from_sheet(ws, start: _date, end: _date, op_col_candidates, source_label: str):
    rows = ws.iter_rows(
        min_row=1, max_row=MAX_ROW_BOUND, min_col=1, max_col=MAX_COL_BOUND, values_only=True
    )
    header_row = next(rows, None)
    if header_row is None:
        return None, 0, 0

    date_col = _find_col(header_row, "date")
    group_col = _find_col(header_row, "group name")
    op_col = _find_col(header_row, *op_col_candidates)

    if date_col is None or op_col is None:
        return "missing_columns", 0, 0

    collected = []
    seen = 0
    for row in rows:
        seen += 1
        raw_date = row[date_col] if date_col < len(row) else None
        if not isinstance(raw_date, (datetime, _date)):
            continue
        d = raw_date.date() if isinstance(raw_date, datetime) else raw_date
        if not (start <= d <= end):
            continue

        op_name = row[op_col] if op_col < len(row) else None
        if op_name is None or str(op_name).strip() == "":
            continue

        group_name = row[group_col] if group_col is not None and group_col < len(row) else None
        collected.append((d, op_name, group_name, source_label))

    return collected, seen, len(collected)


def _load_roster(roster_path: Path):
    """
    Returns {lark_name (stripped): op_name} built from Roster.xlsx's
    first sheet, or raises RuntimeError if the file/columns aren't
    found (Roster is required for the name transfer, so a missing/bad
    Roster is treated as a hard failure, not silently skipped).
    """
    if not roster_path.exists():
        raise RuntimeError(f"Roster file not found: {roster_path}")

    wb = openpyxl.load_workbook(roster_path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(
        min_row=1, max_row=MAX_ROW_BOUND, min_col=1, max_col=MAX_COL_BOUND, values_only=True
    )
    header_row = next(rows, None)
    if header_row is None:
        wb.close()
        raise RuntimeError(f"Roster file has no header row: {roster_path}")

    lark_col = _find_col(header_row, "Lark Name")
    op_col = _find_col(header_row, "OP Name", "CRM OP Name")
    if lark_col is None or op_col is None:
        wb.close()
        raise RuntimeError(
            f"Roster file is missing 'Lark Name' and/or 'OP Name' column: {roster_path}"
        )

    mapping = {}
    for row in rows:
        lark_name = row[lark_col] if lark_col < len(row) else None
        if lark_name is None or str(lark_name).strip() == "":
            continue
        op_name = row[op_col] if op_col < len(row) else None
        key = str(lark_name).strip()
        mapping.setdefault(key, op_name)
    wb.close()
    return mapping


def _apply_roster_transfer(rows, roster_map):
    """
    rows: list of (date, op_name, group_name, source).
    Returns (transferred_rows, matched_count, unmatched_count), where
    transferred_rows is a list of (date, op_name, group_name, source,
    roster_status).
    """
    out = []
    matched = 0
    unmatched = 0
    for d, op_name, group_name, source in rows:
        key = str(op_name).strip() if op_name is not None else ""
        if key in roster_map:
            out.append((d, roster_map[key], group_name, source, None))
            matched += 1
        else:
            out.append((d, op_name, group_name, source, INACTIVE_FLAG))
            unmatched += 1
    return out, matched, unmatched


def combine_target(target: str, start_date: str, end_date: str) -> dict:
    """
    Args:
        target: "ac" or "dw"
        start_date, end_date: "YYYY-MM-DD", inclusive

    Returns a result dict with: target, display_name, expected_files,
    found_files, missing_files, schema_errors, rows, output_path, status.
    Writes the output file only if status is COMPLETE.
    """
    if target not in TARGETS:
        raise ValueError(f"Unknown target '{target}'. Valid: {list(TARGETS)}")
    spec = TARGETS[target]

    start = _date.fromisoformat(start_date)
    end = _date.fromisoformat(end_date)

    result = {
        "target": target,
        "display_name": spec["display_name"],
        "expected_files": len(spec["source_files"]),
        "found_files": 0,
        "missing_files": [],
        "schema_errors": [],
        "rows": 0,
        "roster_matched": 0,
        "roster_unmatched": 0,
        "output_path": None,
        "status": "INCOMPLETE",
    }

    all_rows = []
    for fname in spec["source_files"]:
        fpath = SOURCE_DIR / fname
        if not fpath.exists():
            result["missing_files"].append(fname)
            continue
        result["found_files"] += 1

        file_label = _clean_file_label(fname)

        wb = openpyxl.load_workbook(fpath, read_only=True, data_only=True)
        if spec["sheet_names"]:
            sheet_names = spec["sheet_names"]
        else:
            exclude = set(spec.get("exclude_sheet_names") or [])
            sheet_names = [s for s in wb.sheetnames if s not in exclude]
        for sheet_name in sheet_names:
            if sheet_name not in wb.sheetnames:
                result["schema_errors"].append(
                    {"file": fname, "sheet": sheet_name, "error": "sheet not found"}
                )
                continue
            ws = wb[sheet_name]
            source_label = sheet_name if spec["source_label_mode"] == "sheet_name" else file_label
            collected, seen, kept = _collect_from_sheet(
                ws, start, end, spec["op_col_candidates"], source_label
            )
            if collected == "missing_columns":
                result["schema_errors"].append(
                    {"file": fname, "sheet": sheet_name, "error": "Date/OP-name column not found"}
                )
                continue
            all_rows.extend(collected)
        wb.close()

    if result["missing_files"] or result["schema_errors"]:
        result["status"] = "INCOMPLETE" if result["missing_files"] else "FAILED"
        return result

    try:
        roster_map = _load_roster(ROSTER_PATH)
    except RuntimeError as e:
        result["schema_errors"].append({"file": "Roster.xlsx", "sheet": "-", "error": str(e)})
        result["status"] = "FAILED"
        return result

    transferred_rows, matched, unmatched = _apply_roster_transfer(all_rows, roster_map)
    result["roster_matched"] = matched
    result["roster_unmatched"] = unmatched

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_name = spec["output_name_template"].format(start=start_date, end=end_date)
    output_path = OUTPUT_DIR / output_name

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = spec["display_name"]
    ws_out.append(["Date", "OP Name", "Group Name", "Source", "Roster Status"])
    for d, op_name, group_name, source_label, roster_status in transferred_rows:
        ws_out.append([d, op_name, group_name, source_label, roster_status])
    for r in range(2, len(transferred_rows) + 2):
        ws_out.cell(row=r, column=1).number_format = "yyyy/mm/dd"
    wb_out.save(output_path)
    wb_out.close()

    result["rows"] = len(transferred_rows)
    result["output_path"] = str(output_path)
    result["status"] = "COMPLETE"
    return result


def combine_chatgroup_volumn(start_date: str, end_date: str, targets=None) -> dict:
    """
    Official entry point. Produces the AC and/or DW Chatgroup Volumn
    outputs for the given date range.

    Args:
        start_date, end_date: "YYYY-MM-DD", inclusive
        targets: list of "ac"/"dw" to build; None = both

    Returns {target: result_dict}. Prints a human-readable summary.
    """
    keys = targets if targets else list(TARGETS.keys())
    for key in keys:
        if key not in TARGETS:
            raise ValueError(f"Unknown target '{key}'. Valid: {list(TARGETS)}")

    results = {key: combine_target(key, start_date, end_date) for key in keys}
    _print_summary(start_date, end_date, results)
    return results


def _print_summary(start_date, end_date, results):
    print("Lark Chatgroup Volumn Combine Completed")
    print()
    print("Date Range:")
    print(f"{start_date} -> {end_date}")
    print()
    for result in results.values():
        print(result["display_name"])
        print(f"Expected source files: {result['expected_files']}")
        print(f"Found source files: {result['found_files']}")
        if result["missing_files"]:
            print("Missing files:")
            for f in result["missing_files"]:
                print(f"  {f}")
        if result["schema_errors"]:
            print("Schema errors:")
            for e in result["schema_errors"]:
                print(f"  {e['file']} / {e['sheet']}: {e['error']}")
        print(f"Rows: {result['rows']}")
        if result["rows"]:
            print(
                f"Roster matched: {result['roster_matched']}  "
                f"Roster unmatched (Inactive from Roster): {result['roster_unmatched']}"
            )
        print(f"Status: {result['status']}")
        print()

    outputs = [r["output_path"] for r in results.values() if r["output_path"]]
    if outputs:
        print("Outputs:")
        for o in outputs:
            print(o)
