"""
DW Daily Module Report generator.

Same input data, employee-name matching and date logic as the AC report
(generate_ac_report.py), but counts the 8 DW modules defined in
config.py (DW_*). Rules transcribed from the user's Excel formulas --
see the DW block in config.py.

Employee list source, date range, and output filename are all set in
config.py (DW_*) -- edit those, not this file.

Run:
    python generate_dw_report.py
"""

import random
import sys
from datetime import date, datetime, timedelta

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config
from generate_ac_report import (
    LIGHT_BLUE,
    LIGHT_YELLOW,
    THIN_BORDER,
    build_invalid_data,
    load_employee_names,
    load_input_data,
)

MODULES = config.DW_MODULE_ORDER
GROUP_MODULE_SETS = {
    m: {g.lower() for g in groups} for m, groups in config.DW_GROUP_MODULES.items()
}
OTHER_EXCLUDED = {g.lower() for g in config.DW_OTHER_EXCLUDED_GROUPS}
KEYWORDS = {m: kw.lower() for m, kw in config.DW_KEYWORD_MODULES.items()}


# ---------------------------------------------------------------------------
# Calculation
# ---------------------------------------------------------------------------

def prepare_calc_frame(raw: pd.DataFrame, start_date: date, end_date: date) -> pd.DataFrame:
    df = raw.copy()
    df["parsed_date"] = pd.to_datetime(df["create_time"], errors="coerce").dt.date
    df = df[df["parsed_date"].notna()]
    df = df[(df["parsed_date"] >= start_date) & (df["parsed_date"] <= end_date)]

    group_lower = df["group"].str.lower()
    tags_lower = df["tags"].str.lower()
    for module, groups in GROUP_MODULE_SETS.items():
        df[f"m::{module}"] = group_lower.isin(groups)
    df["m::Other Fresh"] = ~group_lower.isin(OTHER_EXCLUDED)
    for module, kw in KEYWORDS.items():
        df[f"m::{module}"] = tags_lower.str.contains(kw, regex=False, na=False)
    return df


def compute_main_table(df_calc: pd.DataFrame, employee_names, date_range) -> pd.DataFrame:
    flag_cols = [f"m::{m}" for m in MODULES]
    frames = []
    for name in employee_names:
        sub = df_calc[df_calc["tags"].str.contains(name, case=False, regex=False, na=False)]
        counts = sub.groupby("parsed_date")[flag_cols].sum()
        counts = counts.reindex(index=date_range, fill_value=0).fillna(0).astype(int)
        counts.columns = MODULES
        counts.index.name = "Date"
        counts = counts.reset_index()
        counts.insert(1, "OP Name", name)
        frames.append(counts)
    table = pd.concat(frames, ignore_index=True)
    # Row order: by date, then by employee-list order (same as the AC report).
    order = {n: i for i, n in enumerate(employee_names)}
    table["_o"] = table["OP Name"].map(order)
    table = table.sort_values(["Date", "_o"]).drop(columns="_o").reset_index(drop=True)
    return table


def build_other_groups(df_calc: pd.DataFrame) -> pd.DataFrame:
    other = df_calc[df_calc["m::Other Fresh"]]
    counts = other["group"].replace("", "(blank)").value_counts().reset_index()
    counts.columns = ["Group", "Count"]
    return counts


# ---------------------------------------------------------------------------
# Excel writing (same look as the AC report)
# ---------------------------------------------------------------------------

def write_report(main_table, raw, df_calc, other_groups, invalid_data, output_path, sheet_name):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name

    header_fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
    data_fill = PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left_center = Alignment(horizontal="left", vertical="center")
    n_cols = 2 + len(MODULES)

    ws.append([None, "Name"] + MODULES)
    ws.append(["Date", "OP Name"] + [0] * len(MODULES))
    for row in (1, 2):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = header_fill if row == 1 else data_fill
            cell.alignment = center
            cell.border = THIN_BORDER
            cell.font = Font(size=14, bold=True) if (row == 1 and col == 2) else Font(color="000000")

    for r in main_table.itertuples(index=False):
        ws.append(list(r))
    last_row = 2 + len(main_table)
    for row in ws.iter_rows(min_row=3, max_row=last_row, max_col=n_cols):
        for cell in row:
            cell.fill = data_fill
            cell.border = THIN_BORDER
            if cell.column == 1:
                cell.number_format = "yyyy/m/d"
                cell.alignment = center
            elif cell.column == 2:
                cell.alignment = left_center
            else:
                cell.number_format = "0"
                cell.alignment = center

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}{last_row}"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = max(main_table["OP Name"].str.len().max() + 4, 12)
    for i, m in enumerate(MODULES, start=3):
        ws.column_dimensions[get_column_letter(i)].width = max(len(m) + 4, 10)

    bold = Font(bold=True)
    s = wb.create_sheet("Summary")
    s.append(["Metric", "Value"])
    s.append(["Raw data total rows", len(raw)])
    s.append(["Rows within date range", len(df_calc)])
    s.append([])
    s.append(["Module Totals"])
    s.append(["Module", "Total"])
    for m in MODULES:
        s.append([m, int(main_table[m].sum())])
    s.append([])
    s.append(["Employee Totals (all modules combined)"])
    s.append(["OP Name", "Total"])
    emp = main_table.groupby("OP Name", sort=False)[MODULES].sum().sum(axis=1)
    for n, v in emp.items():
        s.append([n, int(v)])
    s.append([])
    s.append(["Daily Totals (all employees / modules combined)"])
    s.append(["Date", "Total"])
    for d, v in main_table.groupby("Date")[MODULES].sum().sum(axis=1).items():
        s.append([d, int(v)])
        s.cell(row=s.max_row, column=1).number_format = "yyyy/m/d"
    for row in s.iter_rows():
        if row[0].value in ("Metric", "Module Totals", "Module", "OP Name", "Date",
                            "Employee Totals (all modules combined)",
                            "Daily Totals (all employees / modules combined)"):
            for c in row:
                c.font = bold
    s.column_dimensions["A"].width = 44
    s.column_dimensions["B"].width = 14

    o = wb.create_sheet("Other Fresh Groups")
    o.append(["Group (counted as Other Fresh)", "Rows in date range"])
    for c in o[1]:
        c.font = bold
    for g, c in other_groups.itertuples(index=False):
        o.append([g, int(c)])
    o.column_dimensions["A"].width = 40
    o.column_dimensions["B"].width = 20

    inv = wb.create_sheet("Invalid Data")
    inv.append(["Row Number", "Created Time", "Group", "Tags", "Reason"])
    for c in inv[1]:
        c.font = bold
    for r in invalid_data.itertuples(index=False):
        inv.append(list(r))
    for col, w in zip("ABCDE", (12, 20, 30, 60, 24)):
        inv.column_dimensions[col].width = w

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def recompute_cell_bruteforce(raw: pd.DataFrame, target_date: date, employee: str, module: str) -> int:
    """Row-by-row re-implementation of the user's Excel formula for one cell."""
    parsed = pd.to_datetime(raw["create_time"], errors="coerce").dt.date
    count = 0
    for _, r in raw[parsed == target_date].iterrows():
        tags, group = r["tags"].lower(), r["group"].lower()
        if employee.lower() not in tags:
            continue
        if module in KEYWORDS:
            hit = KEYWORDS[module] in tags
        elif module == "Other Fresh":
            hit = group not in OTHER_EXCLUDED
        else:
            hit = group in GROUP_MODULE_SETS[module]
        count += int(hit)
    return count


def run_validation(raw, main_table, employee_names, date_range) -> list:
    failures = []
    if set(main_table["Date"]) != set(date_range):
        failures.append("[Check 1] Date coverage mismatch")
    for d, names in main_table.groupby("Date")["OP Name"].apply(set).items():
        if names != set(employee_names):
            failures.append(f"[Check 2] {d} does not have the full employee list")
            break
    for m in MODULES:
        if (main_table[m] < 0).any():
            failures.append(f"[Check 3] Module '{m}' has a negative value")
        if not pd.api.types.is_integer_dtype(main_table[m]):
            failures.append(f"[Check 3] Module '{m}' is not integer-typed")
    if main_table.isna().any().any():
        failures.append("[Check 4] Blank/NaN values found in main table")
    if len(main_table) != len(date_range) * len(employee_names):
        failures.append("[Check 5] Row count != days x employees")

    for m in MODULES:
        cand = main_table[["Date", "OP Name", m]]
        nonzero = cand[cand[m] > 0]
        pool = nonzero if len(nonzero) >= 3 else cand
        for _, r in pool.sample(n=min(3, len(pool)), random_state=42).iterrows():
            actual = recompute_cell_bruteforce(raw, r["Date"], r["OP Name"], m)
            if actual != int(r[m]):
                failures.append(
                    f"[Check 6] Spot-check FAILED: {m} / {r['Date']} / {r['OP Name']}: "
                    f"report={int(r[m])}, recomputed={actual}"
                )
    return failures


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def generate_dw_report(employee_names, output_filename, sheet_name, start_date, end_date):
    """
    Runs the full DW pipeline (load -> calculate -> write -> validate).
    start_date/end_date: "YYYY-MM-DD", inclusive.
    Returns (output_path, ok, module_totals).
    """
    employee_names = list(dict.fromkeys(employee_names))
    print(f"  {len(employee_names)} employees in this report.", flush=True)

    raw = load_input_data()
    print(f"  {len(raw)} raw rows loaded.", flush=True)

    start = datetime.strptime(start_date, "%Y-%m-%d").date()
    end = datetime.strptime(end_date, "%Y-%m-%d").date()
    date_range = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    print(f"  Date range: {start} to {end} ({len(date_range)} days).", flush=True)

    invalid_data = build_invalid_data(raw)
    df_calc = prepare_calc_frame(raw, start, end)
    print(f"  {len(df_calc)} rows within date range.", flush=True)

    main_table = compute_main_table(df_calc, employee_names, date_range)
    other_groups = build_other_groups(df_calc)

    output_path = config.OUTPUT_FOLDER / output_filename
    write_report(main_table, raw, df_calc, other_groups, invalid_data, output_path, sheet_name)

    failures = run_validation(raw, main_table, employee_names, date_range)
    totals = {m: int(main_table[m].sum()) for m in MODULES}
    print()
    if failures:
        print("VALIDATION FAILED:")
        for f in failures:
            print(f"  - {f}")
        return output_path, False, totals
    print("Validation passed: all checks OK.")
    print()
    print("Module totals:")
    for m, v in totals.items():
        print(f"  {m}: {v}")
    print(f"\nOutput file: {output_path.resolve()}")
    return output_path, True, totals


def main():
    employee_names = load_employee_names(list_file=config.DW_EMPLOYEE_LIST_FILE)
    _, ok, _ = generate_dw_report(
        employee_names,
        config.DW_OUTPUT_FILENAME,
        config.DW_MAIN_SHEET_NAME,
        config.DW_START_DATE,
        config.DW_END_DATE,
    )
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
