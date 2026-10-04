"""
AC Daily Module Report generator.

Reads Fresh ticket export data from input/, counts tickets per
day x employee x work module according to the rules in config.py, and
writes a fully-formatted Excel report to output/. Built from scratch --
does not read, depend on, or copy formulas from any pre-existing Excel
template.

Run:
    python generate_ac_report.py
"""

import sys
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import config

LIGHT_BLUE = "ADD8E6"
LIGHT_YELLOW = "FFF9C4"
THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)


# ---------------------------------------------------------------------------
# Loading
# ---------------------------------------------------------------------------

def load_employee_names(list_file=None):
    list_file = list_file if list_file is not None else config.EMPLOYEE_LIST_FILE
    if list_file:
        path = Path(list_file)
        if not path.exists():
            raise RuntimeError(f"EMPLOYEE_LIST_FILE not found: {path}")
        if path.suffix.lower() == ".csv":
            df = pd.read_csv(path, dtype=str, keep_default_na=False, header=None)
        else:
            df = pd.read_excel(path, dtype=str, header=None)
        names = [str(v).strip() for v in df.iloc[:, 0].dropna().tolist() if str(v).strip()]
        # Drop a header cell like "Name"/"OP Name" if present as the first row.
        if names and names[0].lower() in ("name", "op name", "crm op name", "employee", "employee name"):
            names = names[1:]
        if not names:
            raise RuntimeError(f"EMPLOYEE_LIST_FILE has no names: {path}")
        return names
    return list(config.EMPLOYEE_NAMES)


def load_input_data():
    if not config.INPUT_FOLDER.exists():
        raise RuntimeError(f"Input folder not found: {config.INPUT_FOLDER}")

    files = sorted(config.INPUT_FOLDER.glob("*.csv"))
    if not files:
        raise RuntimeError(f"No .csv files found in {config.INPUT_FOLDER}")

    required = list(config.COLUMN_MAPPING.values())
    frames = []
    for f in files:
        df = pd.read_csv(f, dtype=str, keep_default_na=False)
        missing = [c for c in required if c not in df.columns]
        if missing:
            raise RuntimeError(
                f"{f.name} is missing required column(s) {missing}. "
                f"Found columns: {df.columns.tolist()}. "
                f"Fix config.COLUMN_MAPPING to match, or fix the source file."
            )
        frames.append(df[required])

    combined = pd.concat(frames, ignore_index=True)
    combined.columns = ["create_time", "group", "tags"]
    combined.insert(0, "row_number", range(1, len(combined) + 1))
    return combined


# ---------------------------------------------------------------------------
# Calculation
# ---------------------------------------------------------------------------

def classify_group(group: str):
    """Returns the group-based module name, or None if the group does
    not belong to any group-based module (blank group)."""
    if group == "":
        return None
    if group in AC_GROUPS_SET:
        return "AC"
    if group in GROUP_MODULE_INV:
        return GROUP_MODULE_INV[group]
    if group in OTHER2_GROUPS_SET:
        return "Other 2"
    return "Other"


AC_GROUPS_SET = set(config.AC_GROUPS)
GROUP_MODULE_INV = {
    group_name: module_name
    for module_name, group_names in config.GROUP_MODULE_MAP.items()
    for group_name in group_names
}
OTHER2_GROUPS_SET = set(config.OTHER2_GROUPS)
KNOWN_GROUPS_SET = set(config.KNOWN_GROUPS)

GROUP_MODULES = ["AC"] + list(config.GROUP_MODULE_MAP.keys()) + ["Other", "Other 2"]
KEYWORD_MODULES = list(config.KEYWORD_RULES.keys())


def build_invalid_data(raw: pd.DataFrame) -> pd.DataFrame:
    records = []

    parsed = pd.to_datetime(raw["create_time"], errors="coerce")
    for _, r in raw[parsed.isna()].iterrows():
        records.append({
            "Row Number": r["row_number"],
            "Created Time": r["create_time"],
            "Group": r["group"],
            "Tags": r["tags"],
            "Reason": "Date could not be parsed",
        })

    for _, r in raw[raw["tags"].str.strip() == ""].iterrows():
        records.append({
            "Row Number": r["row_number"],
            "Created Time": r["create_time"],
            "Group": r["group"],
            "Tags": r["tags"],
            "Reason": "Tags is blank",
        })

    for _, r in raw[raw["group"].str.strip() == ""].iterrows():
        records.append({
            "Row Number": r["row_number"],
            "Created Time": r["create_time"],
            "Group": r["group"],
            "Tags": r["tags"],
            "Reason": "Group is blank",
        })

    if not records:
        return pd.DataFrame(columns=["Row Number", "Created Time", "Group", "Tags", "Reason"])
    return pd.DataFrame(records).sort_values("Row Number").reset_index(drop=True)


def prepare_calc_frame(raw: pd.DataFrame, start_date: date, end_date: date) -> pd.DataFrame:
    df = raw.copy()
    df["parsed_date"] = pd.to_datetime(df["create_time"], errors="coerce").dt.date
    df = df[df["parsed_date"].notna()]
    df = df[(df["parsed_date"] >= start_date) & (df["parsed_date"] <= end_date)]

    df["group_module"] = df["group"].apply(classify_group)
    tags_lower = df["tags"].str.lower()
    for module_name, keyword in config.KEYWORD_RULES.items():
        df[f"kw::{module_name}"] = tags_lower.str.contains(keyword.lower(), regex=False, na=False)

    return df


def compute_employee_module_counts(df_calc: pd.DataFrame, employee_names, date_range):
    """
    Returns {employee_name: DataFrame indexed by date, columns =
    config.MODULE_ORDER, integer counts}.
    """
    keyword_cols = [f"kw::{m}" for m in KEYWORD_MODULES]
    results = {}

    for name in employee_names:
        mask = df_calc["tags"].str.contains(name, case=False, regex=False, na=False)
        sub = df_calc[mask]

        if len(sub) == 0:
            grp_counts = pd.DataFrame(columns=GROUP_MODULES)
            kw_counts = pd.DataFrame(columns=keyword_cols)
        else:
            grp_counts = (
                sub.groupby(["parsed_date", "group_module"])
                .size()
                .unstack(fill_value=0)
            )
            grp_counts = grp_counts.reindex(columns=GROUP_MODULES, fill_value=0)

            kw_counts = sub.groupby("parsed_date")[keyword_cols].sum()

        combined = pd.concat([grp_counts, kw_counts], axis=1)
        combined = combined.reindex(index=date_range, fill_value=0)
        combined = combined.reindex(columns=GROUP_MODULES + keyword_cols, fill_value=0)
        combined.columns = GROUP_MODULES + KEYWORD_MODULES
        combined = combined[config.MODULE_ORDER].fillna(0).astype(int)

        results[name] = combined

    return results


def build_main_table(employee_names, date_range, employee_frames) -> pd.DataFrame:
    rows = []
    for d in date_range:
        for name in employee_names:
            row = {"Date": d, "OP Name": name}
            counts = employee_frames[name].loc[d]
            for m in config.MODULE_ORDER:
                row[m] = int(counts[m])
            rows.append(row)
    return pd.DataFrame(rows)


def build_other_groups(df_calc: pd.DataFrame) -> pd.DataFrame:
    other_rows = df_calc[df_calc["group_module"] == "Other"]
    if len(other_rows) == 0:
        return pd.DataFrame(columns=["Group", "Count"])
    counts = other_rows["group"].value_counts().reset_index()
    counts.columns = ["Group", "Count"]
    return counts


def build_summary(raw: pd.DataFrame, df_calc: pd.DataFrame, main_table: pd.DataFrame) -> dict:
    module_totals = {m: int(main_table[m].sum()) for m in config.MODULE_ORDER}
    employee_totals = (
        main_table.groupby("OP Name")[config.MODULE_ORDER]
        .sum()
        .sum(axis=1)
        .reindex(config.EMPLOYEE_NAMES if not config.EMPLOYEE_LIST_FILE else None)
    )
    if employee_totals.isna().any():
        employee_totals = main_table.groupby("OP Name")[config.MODULE_ORDER].sum().sum(axis=1)
    daily_totals = main_table.groupby("Date")[config.MODULE_ORDER].sum().sum(axis=1)

    return {
        "raw_row_count": len(raw),
        "rows_in_date_range": len(df_calc),
        "module_totals": module_totals,
        "employee_totals": employee_totals,
        "daily_totals": daily_totals,
    }


# ---------------------------------------------------------------------------
# Excel writing
# ---------------------------------------------------------------------------

def write_report(main_table, summary, other_groups, invalid_data, output_path, sheet_name):
    wb = openpyxl.Workbook()

    _write_main_sheet(wb, main_table, sheet_name)
    _write_summary_sheet(wb, summary)
    _write_other_groups_sheet(wb, other_groups)
    _write_invalid_data_sheet(wb, invalid_data)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def _write_main_sheet(wb, main_table: pd.DataFrame, sheet_name: str):
    ws = wb.active
    ws.title = sheet_name

    header_fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
    data_fill = PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left_center = Alignment(horizontal="left", vertical="center")
    black_font = Font(color="000000")
    name_font = Font(color="000000", size=14, bold=True)

    n_cols = 2 + len(config.MODULE_ORDER)

    # Row 1: module group headers
    row1 = [None, "Name"] + config.MODULE_ORDER
    ws.append(row1)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.alignment = center
        cell.font = name_font if col == 2 else black_font
        cell.border = THIN_BORDER

    # Row 2: reference/label row (static, no formulas)
    row2 = ["Date", "OP Name"] + [0] * len(config.MODULE_ORDER)
    ws.append(row2)
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=2, column=col)
        cell.fill = data_fill
        cell.alignment = center
        cell.font = black_font
        cell.border = THIN_BORDER

    # Row 3+: data
    for _, r in main_table.iterrows():
        ws.append([r["Date"]] + [r["OP Name"]] + [int(r[m]) for m in config.MODULE_ORDER])

    last_row = 2 + len(main_table)
    for row in range(3, last_row + 1):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = data_fill
            cell.border = THIN_BORDER
            if col == 1:
                cell.number_format = "yyyy/m/d"
                cell.alignment = center
            elif col == 2:
                cell.alignment = left_center
            else:
                cell.number_format = "0"
                cell.alignment = center

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}{last_row}"

    # Column widths
    ws.column_dimensions["A"].width = 12
    max_name_len = max([len(n) for n in main_table["OP Name"]] + [len("OP Name")])
    ws.column_dimensions["B"].width = max(max_name_len + 4, 12)
    for i, module in enumerate(config.MODULE_ORDER, start=3):
        col_letter = get_column_letter(i)
        ws.column_dimensions[col_letter].width = max(len(module) + 4, 10)


def _write_summary_sheet(wb, summary: dict):
    ws = wb.create_sheet("Summary")
    bold = Font(bold=True)

    ws.append(["Metric", "Value"])
    for c in ws[1]:
        c.font = bold
    ws.append(["Raw data total rows", summary["raw_row_count"]])
    ws.append(["Rows within date range", summary["rows_in_date_range"]])
    ws.append([])

    ws.append(["Module Totals"])
    ws[ws.max_row][0].font = bold
    ws.append(["Module", "Total"])
    for c in ws[ws.max_row]:
        c.font = bold
    for module, total in summary["module_totals"].items():
        ws.append([module, total])
    ws.append([])

    ws.append(["Employee Totals (all modules combined)"])
    ws[ws.max_row][0].font = bold
    ws.append(["OP Name", "Total"])
    for c in ws[ws.max_row]:
        c.font = bold
    for name, total in summary["employee_totals"].items():
        ws.append([name, int(total)])
    ws.append([])

    ws.append(["Daily Totals (all employees / modules combined)"])
    ws[ws.max_row][0].font = bold
    ws.append(["Date", "Total"])
    for c in ws[ws.max_row]:
        c.font = bold
    for d, total in summary["daily_totals"].items():
        ws.append([d, int(total)])
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy/m/d"

    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 14


def _write_other_groups_sheet(wb, other_groups: pd.DataFrame):
    ws = wb.create_sheet("Other Groups")
    ws.append(["Group", "Count"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for _, r in other_groups.iterrows():
        ws.append([r["Group"], int(r["Count"])])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 12


def _write_invalid_data_sheet(wb, invalid_data: pd.DataFrame):
    ws = wb.create_sheet("Invalid Data")
    ws.append(["Row Number", "Created Time", "Group", "Tags", "Reason"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for _, r in invalid_data.iterrows():
        ws.append([r["Row Number"], r["Created Time"], r["Group"], r["Tags"], r["Reason"]])
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 60
    ws.column_dimensions["E"].width = 24


# ---------------------------------------------------------------------------
# Validation
# ---------------------------------------------------------------------------

def recompute_cell_bruteforce(raw: pd.DataFrame, target_date: date, employee: str, module: str) -> int:
    parsed = pd.to_datetime(raw["create_time"], errors="coerce").dt.date
    day_rows = raw[parsed == target_date]

    count = 0
    for _, r in day_rows.iterrows():
        tags = r["tags"]
        if employee.lower() not in tags.lower():
            continue
        group = r["group"]

        if module in KEYWORD_MODULES:
            keyword = config.KEYWORD_RULES[module]
            if keyword.lower() in tags.lower():
                count += 1
            continue

        if module == "AC":
            if group in AC_GROUPS_SET:
                count += 1
        elif module == "Other":
            if group != "" and group not in KNOWN_GROUPS_SET:
                count += 1
        elif module == "Other 2":
            if group in OTHER2_GROUPS_SET:
                count += 1
        else:
            if group in config.GROUP_MODULE_MAP.get(module, []):
                count += 1

    return count


def run_validation(raw, df_calc, main_table, employee_names, date_range) -> list:
    failures = []

    covered_dates = set(main_table["Date"].unique())
    expected_dates = set(date_range)
    if covered_dates != expected_dates:
        failures.append(
            f"[Check 1] Date coverage mismatch: expected {len(expected_dates)} dates, "
            f"got {len(covered_dates)}"
        )

    grouped = main_table.groupby("Date")["OP Name"].apply(set)
    expected_names = set(employee_names)
    for d, names in grouped.items():
        if names != expected_names:
            failures.append(f"[Check 2] {d} does not have the full employee list")
            break

    for m in config.MODULE_ORDER:
        if (main_table[m] < 0).any():
            failures.append(f"[Check 3] Module '{m}' has a negative value")
        if not pd.api.types.is_integer_dtype(main_table[m]):
            failures.append(f"[Check 3] Module '{m}' is not integer-typed")

    if main_table[config.MODULE_ORDER + ["Date", "OP Name"]].isna().any().any():
        failures.append("[Check 4] Blank/NaN values found in main table")

    expected_rows = len(date_range) * len(employee_names)
    if len(main_table) != expected_rows:
        failures.append(
            f"[Check 5] Row count mismatch: expected {expected_rows} "
            f"({len(date_range)} dates x {len(employee_names)} employees), "
            f"got {len(main_table)}"
        )

    import random
    rng = random.Random(42)
    for module in config.MODULE_ORDER:
        candidates = main_table[["Date", "OP Name", module]]
        nonzero = candidates[candidates[module] > 0]
        pool = nonzero if len(nonzero) >= 3 else candidates
        sample = pool.sample(n=min(3, len(pool)), random_state=42) if len(pool) > 0 else pool

        for _, r in sample.iterrows():
            expected = int(r[module])
            actual = recompute_cell_bruteforce(raw, r["Date"], r["OP Name"], module)
            if actual != expected:
                failures.append(
                    f"[Check 6/7] Spot-check FAILED for module '{module}', "
                    f"date={r['Date']}, employee={r['OP Name']}: "
                    f"report={expected}, recomputed={actual}"
                )

    return failures


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def generate_report(employee_names, output_filename, sheet_name, start_date=None, end_date=None):
    """
    Runs the full pipeline (load -> calculate -> write -> validate) for
    the given employee list, using the shared input folder and module
    rules from config.py. Used by the main 55-employee report and any
    other independent employee-subset/date-range report that needs the
    exact same logic and output format.

    start_date/end_date ("YYYY-MM-DD", inclusive) default to
    config.START_DATE/config.END_DATE when omitted.
    """
    print(f"  {len(employee_names)} employees in this report.", flush=True)

    print("[1/7] Loading input data...", flush=True)
    raw = load_input_data()
    print(f"  {len(raw)} raw rows loaded.", flush=True)

    start_date = datetime.strptime(start_date or config.START_DATE, "%Y-%m-%d").date()
    end_date = datetime.strptime(end_date or config.END_DATE, "%Y-%m-%d").date()
    date_range = [start_date + timedelta(days=i) for i in range((end_date - start_date).days + 1)]
    print(f"  Date range: {start_date} to {end_date} ({len(date_range)} days).", flush=True)

    print("[2/7] Checking for invalid rows...", flush=True)
    invalid_data = build_invalid_data(raw)
    print(f"  {len(invalid_data)} invalid-data entries found.", flush=True)

    print("[3/7] Preparing calculation frame...", flush=True)
    df_calc = prepare_calc_frame(raw, start_date, end_date)
    print(f"  {len(df_calc)} rows within date range.", flush=True)

    print("[4/7] Computing per-employee module counts...", flush=True)
    employee_frames = compute_employee_module_counts(df_calc, employee_names, date_range)

    print("[5/7] Building main table...", flush=True)
    main_table = build_main_table(employee_names, date_range, employee_frames)
    other_groups = build_other_groups(df_calc)
    summary = build_summary(raw, df_calc, main_table)

    print("[6/7] Writing Excel report...", flush=True)
    output_path = config.OUTPUT_FOLDER / output_filename
    write_report(main_table, summary, other_groups, invalid_data, output_path, sheet_name)

    print("[7/7] Running validation...", flush=True)
    failures = run_validation(raw, df_calc, main_table, employee_names, date_range)

    print()
    if failures:
        print("VALIDATION FAILED:")
        for f in failures:
            print(f"  - {f}")
        return output_path, False

    print("Validation passed: all checks OK.")
    print()
    print(f"Output file: {output_path.resolve()}")
    return output_path, True


def main():
    print("[0/7] Loading employee list...", flush=True)
    employee_names = load_employee_names()
    _, ok = generate_report(employee_names, config.OUTPUT_FILENAME, config.MAIN_SHEET_NAME)
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
