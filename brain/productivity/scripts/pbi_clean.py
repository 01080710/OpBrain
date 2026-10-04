#!/usr/bin/env python3
"""
ACT-03  清理 PBI 原始檔

職責：讀每天的原始檔，只留下資料列：去掉第 1 列之外的額外表頭列（R-005）、
      去掉檔尾固定的 3 列非資料（R-004，依位置判斷，不比對文字）。
      原始檔本身不修改。

文件對應：sources.md A（清理與合併）R-004、R-005
輸入：<workdir>/pbi_input_check.json（ACT-02 的輸出）
輸出：<workdir>/pbi_clean_<報表>.pkl  {日期: (表頭[不含 A 欄], 資料列[含 A 欄])}
結束碼：0 成功   1 前一步輸出不存在

用法：python pbi_clean.py --workdir <dir>
"""
import argparse
from datetime import date as _date
from pathlib import Path

import openpyxl

import _bootstrap  # noqa: F401
from opbrain.common import add_workdir, read_json, require, run, save_obj
from pbi_common import cfg


def read_raw(path: Path, header_rows_in_raw: int, trailing: int):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    data_start = header_rows_in_raw + 1
    data_end = ws.max_row - trailing
    rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
            for r in range(data_start, data_end + 1)]
    wb.close()
    return header, rows


def main() -> int:
    ap = argparse.ArgumentParser(description="清理 PBI 原始檔")
    add_workdir(ap)
    args = ap.parse_args()

    check = read_json(require(args.workdir / "pbi_input_check.json", "先執行 ACT-02"))
    c = cfg()
    for k, info in check["reports"].items():
        spec = c["reports"][k]
        cleaned = {}
        for d, f in sorted(info["files"].items()):
            header, rows = read_raw(Path(f), spec["header_rows_in_raw"], c["trailing_non_data_rows"])
            cleaned[_date.fromisoformat(d)] = (header[1:], rows)
        save_obj(args.workdir / f"pbi_clean_{k}.pkl", cleaned)
        print(f"{spec['display_name']}: {len(cleaned)} day(s), "
              f"{sum(len(r) for _, r in cleaned.values())} data row(s)")
    return 0


if __name__ == "__main__":
    run(main)
