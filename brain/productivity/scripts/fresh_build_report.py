#!/usr/bin/env python3
"""
ACT-14  寫出 Fresh AC / DW 報表

職責：把 ACT-13 的結果寫成格式化 Excel 到 data/reports/，4 張工作表：
      主表（Date × OP Name × 模組；淺藍表頭、淺黃資料、凍結窗格、篩選）、Summary、
      Other 清單（AC：Other Groups；DW：Other Fresh Groups）、Invalid Data。

文件對應：flows/ac.md R-046、flows/dw.md R-057
輸入：--kind ac|dw；<workdir>/fresh_<kind>.pkl
輸出：data/reports/<AC|DW> Fresh Report <start> - <end>.xlsx；<workdir>/fresh_<kind>_output.json
結束碼：0 成功   1 前一步輸出不存在 / 檔案被 Excel 開著

用法：python fresh_build_report.py --workdir <dir> --kind dw
"""
import argparse

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import _bootstrap  # noqa: F401
from fresh_common import kind_cfg, output_name
from opbrain import paths
from opbrain.common import add_workdir, load_obj, run, write_json

LIGHT_BLUE, LIGHT_YELLOW = "ADD8E6", "FFF9C4"
THIN = Border(left=Side(style="thin"), right=Side(style="thin"),
              top=Side(style="thin"), bottom=Side(style="thin"))
OTHER_SHEET = {"ac": ("Other Groups", ["Group", "Count"]),
               "dw": ("Other Fresh Groups", ["Group (counted as Other Fresh)", "Rows in date range"])}


def write_main(ws, table, modules, sheet_name):
    ws.title = sheet_name
    head_fill = PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid")
    data_fill = PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid")
    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")
    n_cols = 2 + len(modules)

    ws.append([None, "Name"] + modules)
    ws.append(["Date", "OP Name"] + [0] * len(modules))
    for row in (1, 2):
        for col in range(1, n_cols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = head_fill if row == 1 else data_fill
            cell.alignment = center
            cell.border = THIN
            cell.font = Font(color="000000", size=14, bold=True) if (row, col) == (1, 2) else Font(color="000000")

    for r in table.itertuples(index=False):
        ws.append([r[0], r[1]] + [int(v) for v in r[2:]])
    last = 2 + len(table)
    for row in ws.iter_rows(min_row=3, max_row=last, max_col=n_cols):
        for cell in row:
            cell.fill = data_fill
            cell.border = THIN
            if cell.column == 1:
                cell.number_format, cell.alignment = "yyyy/m/d", center
            elif cell.column == 2:
                cell.alignment = left
            else:
                cell.number_format, cell.alignment = "0", center

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:{get_column_letter(n_cols)}{last}"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = max(max([len(n) for n in table["OP Name"]] + [len("OP Name")]) + 4, 12)
    for i, m in enumerate(modules, start=3):
        ws.column_dimensions[get_column_letter(i)].width = max(len(m) + 4, 10)


def write_summary(wb, res, modules):
    t = res["main_table"]
    ws = wb.create_sheet("Summary")
    bold = Font(bold=True)

    def title(*vals):
        ws.append(list(vals))
        for c in ws[ws.max_row]:
            c.font = bold

    title("Metric", "Value")
    ws.append(["Raw data total rows", res["raw_row_count"]])
    ws.append(["Rows within date range", res["rows_in_range"]])
    ws.append([])
    title("Module Totals")
    title("Module", "Total")
    for m in modules:
        ws.append([m, int(t[m].sum())])
    ws.append([])
    title("Employee Totals (all modules combined)")
    title("OP Name", "Total")
    emp = t.groupby("OP Name", sort=False)[modules].sum().sum(axis=1)   # 依報表名單順序
    for name in res["employees"]:
        ws.append([name, int(emp[name])])
    ws.append([])
    title("Daily Totals (all employees / modules combined)")
    title("Date", "Total")
    for d, v in t.groupby("Date")[modules].sum().sum(axis=1).items():
        ws.append([d, int(v)])
        ws.cell(row=ws.max_row, column=1).number_format = "yyyy/m/d"
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 14


def main() -> int:
    ap = argparse.ArgumentParser(description="寫出 Fresh 報表")
    add_workdir(ap)
    ap.add_argument("--kind", required=True, choices=["ac", "dw"])
    args = ap.parse_args()

    res = load_obj(args.workdir / f"fresh_{args.kind}.pkl")
    modules = list(kind_cfg(args.kind)["module_order"])

    wb = openpyxl.Workbook()
    write_main(wb.active, res["main_table"], modules, kind_cfg(args.kind)["main_sheet_name"])
    write_summary(wb, res, modules)

    sheet, header = OTHER_SHEET[args.kind]
    ws = wb.create_sheet(sheet)
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for g, n in res["other_groups"].itertuples(index=False):
        ws.append([g, int(n)])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 40, 20

    ws = wb.create_sheet("Invalid Data")
    ws.append(list(res["invalid"].columns))
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in res["invalid"].itertuples(index=False):
        ws.append(list(r))
    for col, w in zip("ABCDE", (12, 20, 30, 60, 24)):
        ws.column_dimensions[col].width = w

    paths.REPORTS.mkdir(parents=True, exist_ok=True)
    out = paths.REPORTS / output_name(args.kind, res["employees_preset"], res["start"], res["end"])
    wb.save(out)
    write_json(args.workdir / f"fresh_{args.kind}_output.json", {"output_path": str(out)})
    print(f"{args.kind.upper()} report: {out}")
    return 0


if __name__ == "__main__":
    run(main)
