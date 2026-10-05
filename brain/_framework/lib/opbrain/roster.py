"""Roster 讀取（所有 Skill 共用）。

規則擁有者：team-structure（brain/management/team-structure/rules.md R-001～R-006），
規則的值：brain/management/team-structure/config/roster.yaml。
其他 Skill 不得自行讀 Roster，一律透過這裡，避免各自實作出不同結果。
"""
from pathlib import Path

import openpyxl
import pandas as pd

from . import config, paths
from .common import InputError

_CFG_SKILL = "team-structure"


def _cfg():
    return config.load(_CFG_SKILL, "roster")


def default_path() -> Path:
    return paths.PROJECT_ROOT / _cfg()["roster_file"]  # R-001


def find_col(header_row, *candidates):
    """依表頭名稱找欄位（去空白、不分大小寫）。找不到回傳 None。"""
    normalized = {}
    for i, h in enumerate(header_row):
        if h is None:
            continue
        normalized.setdefault(str(h).strip().lower(), i)
    for cand in candidates:
        idx = normalized.get(cand.strip().lower())
        if idx is not None:
            return idx
    return None


def _rows(path: Path):
    path = Path(path)
    if not path.exists():
        raise InputError(f"Roster file not found: {path}")
    b = _cfg()["read_bounds"]  # R-006
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]  # R-001
    rows = list(ws.iter_rows(min_row=1, max_row=b["max_row"], min_col=1,
                             max_col=b["max_col"], values_only=True))
    wb.close()
    if not rows:
        raise InputError(f"Roster file has no header row: {path}")
    return rows[0], rows[1:]


def op_name_set(path=None) -> set:
    """正式姓名欄的所有非空白姓名（去前後空白）。R-002、R-003。"""
    header, rows = _rows(path or default_path())
    col = find_col(header, *_cfg()["op_name_headers"])
    if col is None:
        raise InputError(f"Roster file is missing 'OP Name' column: {path or default_path()}")
    names = set()
    for row in rows:
        v = row[col] if col < len(row) else None
        if v is not None and str(v).strip():
            names.add(str(v).strip())
    return names


def lark_to_op_map(path=None) -> dict:
    """{Lark Name（去前後空白）: 正式姓名（原值）}，同一 Lark Name 以第一筆為準。R-004。"""
    header, rows = _rows(path or default_path())
    cfg = _cfg()
    lark_col = find_col(header, *cfg["lark_name_headers"])
    op_col = find_col(header, *cfg["op_name_headers"])
    if lark_col is None or op_col is None:
        raise InputError(
            f"Roster file is missing 'Lark Name' and/or 'OP Name' column: {path or default_path()}"
        )
    mapping = {}
    for row in rows:
        lark = row[lark_col] if lark_col < len(row) else None
        if lark is None or str(lark).strip() == "":
            continue
        op = row[op_col] if op_col < len(row) else None
        mapping.setdefault(str(lark).strip(), op)
    return mapping


def employee_list(path=None) -> list:
    """以第一欄為員工名單：略過空白與表頭列、去重並保留順序。R-003、R-005。"""
    path = Path(path or default_path())
    if not path.exists():
        raise InputError(f"employee list file not found: {path}")
    if path.suffix.lower() == ".csv":
        df = pd.read_csv(path, dtype=str, keep_default_na=False, header=None)
    else:
        df = pd.read_excel(path, dtype=str, header=None)
    names = [str(v).strip() for v in df.iloc[:, 0].dropna().tolist() if str(v).strip()]
    if names and names[0].lower() in _cfg()["header_like_first_cells"]:
        names = names[1:]
    if not names:
        raise InputError(f"employee list file has no names: {path}")
    return list(dict.fromkeys(names))


def first_column_is_name_column(path=None) -> bool:
    header, _ = _rows(path or default_path())
    first = str(header[0]).strip().lower() if header and header[0] is not None else ""
    return first in [h.lower() for h in _cfg()["op_name_headers"]]
