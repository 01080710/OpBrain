#!/usr/bin/env python3
"""
ACT-15  驗證 Fresh 報表

職責：重新開啟寫出的報表主表，檢查（V-020～V-026）：
      日期涵蓋完整、每天都有完整員工名單、件數為非負整數、沒有空值、
      列數 = 天數 × 人數，以及每個模組抽 3 格用「逐列重算」（與計算步驟不同的寫法）比對。
      未通過 → 報表改名加上「VALIDATION FAILED」，避免被誤用。

文件對應：sources.md B V-020～V-026、context.md C-001
輸入：--kind ac|dw；<workdir>/fresh_<kind>.pkl、fresh_<kind>_output.json
輸出：<workdir>/fresh_<kind>_validation.json
結束碼：0 全部通過   1 前一步輸出不存在   2 驗證未通過（關卡）

用法：python fresh_validate_output.py --workdir <dir> --kind dw
"""
import argparse
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

import _bootstrap  # noqa: F401
from fresh_common import load_short_data
from fresh_rules import Rules
from opbrain.common import add_workdir, load_obj, read_json, require, run, write_json


def main() -> int:
    ap = argparse.ArgumentParser(description="驗證 Fresh 報表")
    add_workdir(ap)
    ap.add_argument("--kind", required=True, choices=["ac", "dw"])
    args = ap.parse_args()

    res = load_obj(args.workdir / f"fresh_{args.kind}.pkl")
    out = Path(read_json(require(args.workdir / f"fresh_{args.kind}_output.json", "先執行 ACT-14"))["output_path"])
    rules = Rules(args.kind)
    modules = rules.modules

    wb = openpyxl.load_workbook(out, read_only=True)
    rows = list(wb[wb.sheetnames[0]].iter_rows(min_row=3, values_only=True))
    wb.close()
    t = pd.DataFrame(rows, columns=["Date", "OP Name"] + modules)
    t["Date"] = [d.date() if isinstance(d, datetime) else d for d in t["Date"]]

    names, dates = res["employees"], res["date_range"]
    fails = []
    if set(t["Date"]) != set(dates):
        fails.append("V-020 date coverage mismatch")
    for d, got in t.groupby("Date")["OP Name"].apply(set).items():
        if got != set(names):
            fails.append(f"V-021 {d} does not have the full employee list")
            break
    if t[modules + ["Date", "OP Name"]].isna().any().any():
        fails.append("V-023 blank values in main table")
    else:
        for m in modules:
            if not all(isinstance(v, int) and v >= 0 for v in t[m]):
                fails.append(f"V-022 module '{m}' has a non-integer or negative value")
    if len(t) != len(dates) * len(names):
        fails.append(f"V-024 row count {len(t)} != {len(dates)} days x {len(names)} employees")

    raw = load_short_data()
    parsed = pd.to_datetime(raw["create_time"], errors="coerce").dt.date
    for m in modules:
        if fails and any(f.startswith("V-023") or f.startswith("V-022") for f in fails):
            break
        cand = t[["Date", "OP Name", m]]
        nonzero = cand[cand[m] > 0]
        pool = nonzero if len(nonzero) >= 3 else cand
        for _, r in pool.sample(n=min(3, len(pool)), random_state=42).iterrows():
            day = raw[parsed == r["Date"]]
            actual = sum(
                1 for g, tags in zip(day["group"], day["tags"])
                if r["OP Name"].lower() in tags.lower() and rules.brute_force_hit(g, tags, m)
            )
            if actual != int(r[m]):
                fails.append(f"V-025 spot-check {m} / {r['Date']} / {r['OP Name']}: "
                             f"report={int(r[m])}, recomputed={actual}")

    final_path = out
    if fails:
        final_path = out.with_name(out.stem + " - VALIDATION FAILED" + out.suffix)
        out.replace(final_path)
        print("VALIDATION FAILED:")
        for f in fails:
            print(f"  - {f}")
    else:
        print(f"{args.kind.upper()} validation passed: all checks OK ({len(t)} rows)")
    write_json(args.workdir / f"fresh_{args.kind}_validation.json",
               {"ok": not fails, "failures": fails, "output_path": str(final_path)})
    return 2 if fails else 0


if __name__ == "__main__":
    run(main)
