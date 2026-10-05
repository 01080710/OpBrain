"""chatgroup 流程各步驟共用的函式。"""
import openpyxl

from opbrain import config, paths
from opbrain.common import InputError
from opbrain.roster import find_col


def cfg():
    return config.load("productivity", "chatgroup")


def source_path():
    return paths.PROJECT_ROOT / cfg()["source_file"]  # R-060


def target_keys(arg_list):
    keys = arg_list or list(cfg()["targets"])
    for k in keys:
        if k not in cfg()["targets"]:
            raise InputError(f"Unknown target '{k}'. Valid: {list(cfg()['targets'])}")
    return keys


def sheets_for(spec, wb_sheetnames):
    """R-061：DW = 指定的 3 張；AC = 檔案裡其餘全部。"""
    if spec["sheet_names"]:
        return list(spec["sheet_names"])
    exclude = set(spec.get("exclude_sheet_names") or [])
    return [s for s in wb_sheetnames if s not in exclude]


def sheet_rows(ws):
    b = cfg()["read_bounds"]  # R-066
    return ws.iter_rows(min_row=1, max_row=b["max_row"], min_col=1, max_col=b["max_col"], values_only=True)


def header_columns(header_row, spec):
    """R-062：依表頭名稱找 Date / Group Name / OP 欄。"""
    c = cfg()
    return (find_col(header_row, c["date_header"]),
            find_col(header_row, c["group_header"]),
            find_col(header_row, *spec["op_col_candidates"]))


def open_source():
    p = source_path()
    if not p.exists():
        raise InputError(f"source file not found: {p}")
    return openpyxl.load_workbook(p, read_only=True, data_only=True)
